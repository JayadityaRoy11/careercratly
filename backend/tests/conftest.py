from __future__ import annotations

import os
from io import BytesIO
from pathlib import Path

os.environ["ENVIRONMENT"] = "test"
os.environ["DEBUG"] = "false"
os.environ["PRELOAD_EXTRACTOR"] = "false"
os.environ["UPLOAD_DIR"] = str(Path(__file__).resolve().parents[1] / "uploads" / "test")

import pytest
from fastapi.testclient import TestClient

from app.main import app
from app.services.resume_store import resume_store


@pytest.fixture(autouse=True)
def _clear_store() -> None:
    resume_store.clear()
    yield
    resume_store.clear()


@pytest.fixture
def client() -> TestClient:
    return TestClient(app)


def make_resume_docx(text: str) -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("Jane Doe")
    document.add_heading("Education", level=1)
    document.add_paragraph("B.Tech Computer Science")
    document.add_heading("Experience", level=1)
    document.add_paragraph("Software intern. Developed and deployed APIs.")
    document.add_heading("Projects", level=1)
    document.add_paragraph("Built a FastAPI service.")
    document.add_heading("Skills", level=1)
    document.add_paragraph(text)
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def make_resume_pdf(text: str) -> bytes:
    from reportlab.lib.pagesizes import letter
    from reportlab.pdfgen import canvas

    buffer = BytesIO()
    pdf = canvas.Canvas(buffer, pagesize=letter)
    pdf.drawString(72, 720, "Jane Doe")
    pdf.drawString(72, 700, "Education")
    pdf.drawString(72, 684, "B.Tech Computer Science")
    pdf.drawString(72, 660, "Experience")
    pdf.drawString(72, 644, "Software intern. Developed and deployed APIs.")
    pdf.drawString(72, 620, "Projects")
    pdf.drawString(72, 604, "Built a FastAPI service.")
    pdf.drawString(72, 580, "Skills")
    pdf.drawString(72, 564, text)
    pdf.save()
    return buffer.getvalue()


SAMPLE_SKILLS = "Python FastAPI PostgreSQL Docker SQL Git REST APIs pandas numpy"
