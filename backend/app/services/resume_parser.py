"""Resume text extraction.

Swap PDF/DOCX libraries here without changing upload or storage code.
"""

from __future__ import annotations

import io
from pathlib import Path
from typing import Any

import pdfplumber
from docx import Document

SUPPORTED_EXTENSIONS = {".pdf", ".docx"}
PDF_MAGIC = b"%PDF"
DOCX_MAGIC = b"PK\x03\x04"


class ResumeParseError(Exception):
    def __init__(self, message: str, *, code: str = "parse_error") -> None:
        self.code = code
        super().__init__(message)


def detect_extension(filename: str) -> str:
    return Path(filename).suffix.lower()


def extract_text(file_bytes: bytes, filename: str) -> str:
    return parse_resume(file_bytes, filename)["text"]


def parse_resume(file_bytes: bytes, filename: str) -> dict[str, Any]:
    if not file_bytes:
        raise ResumeParseError("Uploaded file is empty", code="empty_file")

    extension = detect_extension(filename)
    if extension not in SUPPORTED_EXTENSIONS:
        raise ResumeParseError(
            "Unsupported file type. Upload a PDF or DOCX file.",
            code="unsupported_type",
        )

    if extension == ".pdf":
        _assert_pdf_magic(file_bytes)
        text, pages, metadata = _extract_pdf(file_bytes)
        file_type = "pdf"
    else:
        text, pages, metadata = _extract_docx(file_bytes)
        file_type = "docx"

    return {
        "filename": Path(filename).name,
        "file_type": file_type,
        "text": text,
        "pages": pages,
        "metadata": metadata,
    }


def _assert_pdf_magic(file_bytes: bytes) -> None:
    if not file_bytes.lstrip().startswith(PDF_MAGIC):
        raise ResumeParseError("File is not a valid PDF.", code="corrupt_file")


def _extract_pdf(file_bytes: bytes) -> tuple[str, int, dict[str, Any]]:
    try:
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages = list(pdf.pages)
            if not pages:
                raise ResumeParseError("PDF has no pages.", code="corrupt_file")
            chunks = [(page.extract_text() or "") for page in pages]
            info = pdf.metadata or {}
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("Corrupt or unreadable PDF.", code="corrupt_file") from exc

    text = "\n".join(chunks).strip()
    if not text:
        raise ResumeParseError(
            "No extractable text found in the PDF.",
            code="empty_text",
        )
    metadata = {
        "parser": "pdfplumber",
        "page_count": len(pages),
        "title": str(info.get("Title") or "") or None,
        "author": str(info.get("Author") or "") or None,
    }
    return text, len(pages), metadata


def _extract_docx(file_bytes: bytes) -> tuple[str, int, dict[str, Any]]:
    if not file_bytes.startswith(DOCX_MAGIC):
        raise ResumeParseError("File is not a valid DOCX.", code="corrupt_file")

    try:
        document = Document(io.BytesIO(file_bytes))
        parts: list[str] = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.append("\t".join(cell.text for cell in row.cells))
        core = document.core_properties
        metadata = {
            "parser": "python-docx",
            "author": core.author or None,
            "title": core.title or None,
        }
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("Corrupt or unreadable DOCX.", code="corrupt_file") from exc

    text = "\n".join(part for part in parts if part and part.strip()).strip()
    if not text:
        raise ResumeParseError(
            "No extractable text found in the DOCX.",
            code="empty_text",
        )
    return text, 0, metadata
